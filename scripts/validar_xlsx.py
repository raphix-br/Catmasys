from openpyxl import load_workbook
from pathlib import Path

arquivo = Path("dados/teste.xlsx")

if not arquivo.exists():
    raise FileNotFoundError(f"Arquivo não encontrado: {arquivo}")

wb = load_workbook(arquivo, read_only=True, data_only=False)

print("=" * 60)
print("VALIDAÇÃO DO TESTE.XLSX")
print("=" * 60)

print(f"\nArquivo: {arquivo}")
print(f"Planilhas: {wb.sheetnames}")

for ws in wb.worksheets:
    print("\n" + "-" * 60)
    print(f"PLANILHA: {ws.title}")
    print(f"Linhas: {ws.max_row}")
    print(f"Colunas: {ws.max_column}")

    headers = []
    for cell in ws[1]:
        headers.append(cell.value)

    print("\nCABEÇALHOS:")
    for i, header in enumerate(headers, 1):
        print(f"{i}: {header}")

    print("\nPRIMEIRAS 5 LINHAS:")
    for row in ws.iter_rows(min_row=2, max_row=min(6, ws.max_row), values_only=True):
        print(row)

print("\n" + "=" * 60)
print("VALIDAÇÃO CONCLUÍDA")
print("=" * 60)
