from pathlib import Path
import csv
import re
import json
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "dados" / "teste.xlsx"
CSV_OUT = ROOT / "dados" / "Catalogo_BR.csv"
MAPA = ROOT / "dados" / "mapa_capas.csv"
JS_OUT = ROOT / "js" / "games_br.js"

EXPECTED = [
    "id","Titulo_BR","Titulo_Original","Codigo_TecToy","Ano","Variante","Edicao",
    "Genero","Sub-Genero","Peripherals","Players","Tamanho","Desenvolvedor",
    "Distribuidor","FM","Observacao","Exclusivo_BR"
]

def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()

def norm(s):
    s = clean(s)
    s = re.sub(r"[^a-zA-Z0-9]+", "_", s).strip("_").lower()
    return s

def main():
    if not XLSX.exists():
        raise FileNotFoundError(XLSX)

    wb = load_workbook(XLSX, read_only=True, data_only=True)
    if "Catalogo_BR" not in wb.sheetnames:
        raise RuntimeError(f"Planilha Catalogo_BR nao encontrada. Abas: {wb.sheetnames}")

    ws = wb["Catalogo_BR"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise RuntimeError("Aba Catalogo_BR vazia.")

    headers = [clean(x) for x in rows[0]]
    idx = {h:i for i,h in enumerate(headers)}

    missing = [h for h in EXPECTED if h not in idx]
    if missing:
        raise RuntimeError("Colunas obrigatorias ausentes: " + ", ".join(missing))

    records = []
    seen = set()
    for raw in rows[1:]:
        vals = {h: clean(raw[idx[h]] if idx[h] < len(raw) else "") for h in EXPECTED}
        if not vals["id"]:
            continue
        if vals["id"] in seen:
            raise RuntimeError(f"ID duplicado na planilha master: {vals['id']}")
        seen.add(vals["id"])
        records.append(vals)

    if not records:
        raise RuntimeError("Nenhum registro valido encontrado.")

    cover = {}
    if MAPA.exists():
        with MAPA.open("r", encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f):
                if row.get("ID"):
                    cover[row["ID"].strip()] = {
                        "image": clean(row.get("Capa")),
                    }

    missing_cover = [r["id"] for r in records if r["id"] not in cover or not cover[r["id"]]["image"]]
    if missing_cover:
        raise RuntimeError("IDs sem capa em mapa_capas.csv: " + ", ".join(missing_cover))

    with CSV_OUT.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=EXPECTED)
        w.writeheader()
        w.writerows(records)

    games = []
    for r in records:
        c = cover[r["id"]]["image"]
        games.append({
            "id": r["id"],
            "name": r["Titulo_BR"],
            "year": r["Ano"],
            "edition": r["Edicao"],
            "image": c,
            "thumb": re.sub(r"\.[^.]+$", ".webp", c.replace("CAPAS_MASTER_SYSTEM/", "thumbs/")),
            "titulo_original": r["Titulo_Original"],
            "codigo_tectoy": r["Codigo_TecToy"],
            "variante": r["Variante"],
            "genero": r["Genero"],
            "sub-genero": r["Sub-Genero"],
            "peripherals": r["Peripherals"],
            "players": r["Players"],
            "tamanho": r["Tamanho"],
            "desenvolvedor": r["Desenvolvedor"],
            "distribuidor": r["Distribuidor"],
            "fm": r["FM"],
            "observacao": r["Observacao"],
            "exclusivo_br": r["Exclusivo_BR"],
        })

    JS_OUT.write_text(
        "window.GAMES = " + json.dumps(games, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8"
    )

    print(f"OK: {len(records)} jogos gerados.")
    print(f"CSV: {CSV_OUT}")
    print(f"JS:  {JS_OUT}")

if __name__ == "__main__":
    main()

# workflow bootstrap
